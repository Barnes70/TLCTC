#!/usr/bin/env python3
"""Extract the FACTUAL skeleton of the CIS Controls v8.1.x from the CIS spreadsheet.

Emits ids, control number, asset class, security function and minimum implementation group
only. CIS titles and descriptions are CC BY-NC-ND 4.0 and never leave the spreadsheet.

Usage: python tools/extract_skeleton.py <CIS_Controls_Version_8.1.2.xlsx> <out.json>
"""
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

COLS = {'control': ('cis control',), 'safeguard': ('cis safeguard',), 'asset_class': ('asset class', 'asset type'),
        'function': ('security function',), 'ig1': ('ig1',), 'ig2': ('ig2',), 'ig3': ('ig3',)}
FUNCS = {'govern': 'GV', 'identify': 'ID', 'protect': 'PR', 'detect': 'DE', 'respond': 'RS', 'recover': 'RC'}


def _clean(v):
    return str(v).replace('\xa0', ' ').strip() if v is not None else ''


def _find_header(wb):
    """The controls sheet is the one whose header row (within its first 20 rows) has "CIS Safeguard"."""
    for ws in wb.worksheets:
        for i, row in enumerate(ws.iter_rows(min_row=1, max_row=20, values_only=True), start=1):
            names = [_clean(v).lower() for v in row]
            if 'cis safeguard' not in names:
                continue
            ix = {}
            for key, options in COLS.items():
                found = next((names.index(o) for o in options if o in names), None)
                if found is None:
                    raise ValueError(f'sheet {ws.title!r} row {i}: no column {options[0]!r}; found {names}')
                ix[key] = found
            return ws, i, ix
    raise ValueError('no sheet with a header row containing "CIS Safeguard" in its first 20 rows')


def extract(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    try:  # read-only workbooks hold the file open until closed (blocks deletion on Windows)
        return _rows(wb)
    finally:
        wb.close()


def _rows(wb):
    ws, head_row, ix = _find_header(wb)
    out, pos = [], {}
    for row in ws.iter_rows(min_row=head_row + 1, values_only=True):
        if len(row) <= max(ix.values()):
            continue
        c, s = _clean(row[ix['control']]), row[ix['safeguard']]
        if s in (None, '') or not c:
            continue  # control header rows and blank rows
        control = int(c)
        pos[control] = pos.get(control, 0) + 1
        expected = f'{control}.{pos[control]}'
        if isinstance(s, str):
            sid = _clean(s)
            if sid != expected:
                raise ValueError(f'safeguard {sid!r} at position {pos[control]} of control {control}, expected {expected}')
        else:  # numeric cell: "4.10" arrives as 4.1, so the id comes from the position and is cross-checked
            if abs(float(s) - float(expected)) > 1e-9:
                raise ValueError(f'numeric safeguard {s!r} at position {pos[control]} of control {control}, expected {expected}')
            sid = expected
        fn = FUNCS.get(_clean(row[ix['function']]).lower())
        if not fn:
            raise ValueError(f'{sid}: unknown security function {row[ix["function"]]!r}')
        ig = next((n for n, k in ((1, 'ig1'), (2, 'ig2'), (3, 'ig3')) if _clean(row[ix[k]])), None)
        if ig is None:
            raise ValueError(f'{sid}: no implementation group marked')
        out.append({'id': sid, 'control': control, 'asset_class': _clean(row[ix['asset_class']]), 'function': fn, 'ig': ig})
    return out


def write_skeleton(rows, out_path):
    meta = {
        'schema': 'tlctc-cis-v8.1-mapping.v1',
        'title': 'CIS Critical Security Controls v8.1.2 Safeguards × TLCTC clusters',
        'tlctc_version': '2.6',
        'source': {'title': 'CIS Critical Security Controls Version 8.1.2', 'date': '2025-03', 'publisher': 'Center for Internet Security',
                   'url': 'https://www.cisecurity.org/controls/v8-1',
                   'license': 'CC BY-NC-ND 4.0 (https://creativecommons.org/licenses/by-nc-nd/4.0/)'},
        'text_policy': 'No CIS text. Ids, asset classes, security functions and implementation groups are facts; topic and rationale are TLCTC wording. Read the Safeguards in your own copy of the CIS Controls.',
        'mapping_date': None,
        'status': 'AI-assisted starter guidance; owner review pending. Not a CIS product.',
        'clusters_field': 'Array of cluster ids = the Safeguard addresses those rows (Local). "all" = cluster-neutral, shared Umbrella in every row of its function column. "none" = outside the threat axis.',
        'hardening_field': 'true = the Safeguard sets, restricts or verifies the configuration state of an asset or software.',
    }
    blank = {'topic': '', 'clusters': None, 'kind': None, 'hardening': None, 'rationale': ''}
    lines = [json.dumps({**r, **blank}, ensure_ascii=False) for r in rows]
    text = ('{\n  "meta": ' + json.dumps(meta, ensure_ascii=False) + ',\n  "safeguards": [\n    '
            + ',\n    '.join(lines) + '\n  ]\n}\n')
    Path(out_path).write_bytes(text.replace('\n', '\r\n').encode('utf-8'))


if __name__ == '__main__':
    src, dst = sys.argv[1:3]
    rows = extract(src)
    write_skeleton(rows, dst)
    print(f'{len(rows)} safeguards in {len({r["control"] for r in rows})} controls -> {dst}')
