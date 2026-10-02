#!/usr/bin/env python3
"""Fail if any topic/rationale shares a run of >= 6 consecutive words with CIS text in the spreadsheet.

Usage: python tools/check_no_cis_text.py <CIS xlsx> <mapping.json>
"""
import json
import re
import sys

from openpyxl import load_workbook

WORD = re.compile(r'[a-z0-9]+')


def _grams(text, n):
    w = WORD.findall(text.lower())
    return {' '.join(w[i:i + n]) for i in range(len(w) - n + 1)}


def cis_texts(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        return [v for ws in wb.worksheets for row in ws.iter_rows(values_only=True)
                for v in row if isinstance(v, str) and len(v.split()) >= 6]
    finally:
        wb.close()


def leaks(mapping, texts, n=6):
    pool = set().union(*(_grams(t, n) for t in texts)) if texts else set()
    hits = []
    for s in mapping['safeguards']:
        for field in ('topic', 'rationale'):
            shared = sorted(_grams(s.get(field) or '', n) & pool)
            if shared:
                hits.append((s['id'], field, shared[0]))
    return hits


if __name__ == '__main__':
    xlsx, js = sys.argv[1:3]
    with open(js, encoding='utf-8') as f:
        hits = leaks(json.load(f), cis_texts(xlsx))
    for h in hits:
        print('CIS TEXT:', *h)
    print(f'{len(hits)} overlap(s)')
    sys.exit(1 if hits else 0)
