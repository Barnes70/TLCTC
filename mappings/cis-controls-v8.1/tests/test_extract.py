import json
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'tools'))
from extract_skeleton import extract, write_skeleton  # noqa: E402
from check_no_cis_text import leaks  # noqa: E402

HEAD = ['CIS Control', 'CIS Safeguard', 'Asset Class', 'Security Function', 'Title', 'Description', 'IG1', 'IG2', 'IG3']


def book(rows, tmp):
    """Synthetic workbook shaped like the CIS release: the controls sheet is not the first sheet."""
    wb = Workbook()
    wb.active.title = 'Introduction'
    wb.active.append(['Synthetic introduction text'])
    ws = wb.create_sheet('Controls')
    ws.append(HEAD)
    for r in rows:
        ws.append(r)
    p = Path(tmp) / 'cis.xlsx'
    wb.save(p)
    return p


def ctl(n):  # control header row: number (with a trailing no-break space, as in the release), no safeguard
    return [f'{n}\xa0', None, None, None, 'Control heading text', 'Control overview text', None, None, None]


def sg(c, s, fn='Protect', ig=(None, 'x', 'x')):
    return [str(c), s, 'Software', fn, f'Synthetic title {s}', f'Synthetic description {s}', *ig]


class Extract(unittest.TestCase):
    def test_text_ids_and_skips_control_rows(self):
        with tempfile.TemporaryDirectory() as t:
            rows = extract(book([ctl(4)] + [sg(4, f'4.{i}') for i in range(1, 11)], t))
        self.assertEqual([r['id'] for r in rows], [f'4.{i}' for i in range(1, 11)])
        self.assertEqual(rows[9], {'id': '4.10', 'control': 4, 'asset_class': 'Software', 'function': 'PR', 'ig': 2})

    def test_numeric_ids_recovered_by_position(self):
        nums = [4.1, 4.2, 4.3, 4.4, 4.5, 4.6, 4.7, 4.8, 4.9, 4.1]  # Excel turned "4.10" into 4.1
        with tempfile.TemporaryDirectory() as t:
            rows = extract(book([ctl(4)] + [sg(4, n) for n in nums], t))
        self.assertEqual(rows[-1]['id'], '4.10')
        self.assertEqual(len({r['id'] for r in rows}), 10)

    def test_numeric_id_out_of_position_is_an_error(self):
        with tempfile.TemporaryDirectory() as t:
            p = book([ctl(4), sg(4, 4.1), sg(4, 4.3)], t)
            with self.assertRaisesRegex(ValueError, 'position'):
                extract(p)

    def test_function_and_ig(self):
        with tempfile.TemporaryDirectory() as t:
            rows = extract(book([ctl(1), sg(1, '1.1', 'Identify', ('x', 'x', 'x')), sg(1, '1.2', 'Govern', (None, None, 'x'))], t))
        self.assertEqual([(r['function'], r['ig']) for r in rows], [('ID', 1), ('GV', 3)])

    def test_missing_header_names_found(self):
        wb = Workbook()
        wb.active.append(['nothing', 'here'])
        with tempfile.TemporaryDirectory() as t:
            p = Path(t) / 'x.xlsx'
            wb.save(p)
            with self.assertRaisesRegex(ValueError, 'CIS Safeguard'):
                extract(p)

    def test_skeleton_is_one_line_per_safeguard_crlf(self):
        rows = [{'id': '4.1', 'control': 4, 'asset_class': 'Software', 'function': 'PR', 'ig': 1}]
        with tempfile.TemporaryDirectory() as t:
            out = Path(t) / 'm.json'
            write_skeleton(rows, out)
            raw = out.read_bytes().decode('utf-8')
            self.assertIn('\r\n', raw)
            line = [ln for ln in raw.split('\r\n') if '"4.1"' in ln][0]
            self.assertIsNone(json.loads(line.strip().rstrip(','))['clusters'])
            self.assertEqual(len(json.loads(raw)['safeguards']), 1)


class Leaks(unittest.TestCase):
    def test_six_word_overlap_flagged_five_not(self):
        texts = ['Ensure that unnecessary services are uninstalled or disabled on enterprise assets']
        m = {'safeguards': [
            {'id': '4.8', 'topic': 'unused services off', 'rationale': 'services are uninstalled or disabled on enterprise assets here'},
            {'id': '4.9', 'topic': 'x', 'rationale': 'or disabled on enterprise assets'}]}  # 5 shared words only
        hits = leaks(m, texts)
        self.assertEqual([(h[0], h[1]) for h in hits], [('4.8', 'rationale')])


if __name__ == '__main__':
    unittest.main()
